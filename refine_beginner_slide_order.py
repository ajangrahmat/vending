"""Place the guided slides in teaching order and fix their step-number labels."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


PATH = r"D:\workshop-sukabumi\vending\materi_praktik_vending_90_menit.pptx"
deck = Presentation(PATH)
slide_ids = list(deck.slides._sldIdLst)
if len(slide_ids) != 27:
    raise ValueError(f"Expected 27 slides before ordering; got {len(slide_ids)}.")

original = slide_ids[:19]
guided = slide_ids[19:]
ordered = (
    original[:5]
    + guided[:3]
    + original[5:6]
    + guided[3:5]
    + original[6:7]
    + guided[5:6]
    + original[7:15]
    + [guided[7]]
    + original[15:16]
    + [guided[6]]
    + original[16:]
)
slide_id_list = deck.slides._sldIdLst
for slide_id in list(slide_id_list):
    slide_id_list.remove(slide_id)
for slide_id in ordered:
    slide_id_list.append(slide_id)

agenda = deck.slides[2]
agenda_labels = [
    "00–08  Orientasi dan target belajar",
    "08–18  Buka proyek dan jalankan aplikasi",
    "18–28  Coba aplikasi dalam simulasi",
    "28–48  Ubah produk dan periksa hasil",
    "48–73  Telusuri kode dan tes serial",
    "73–90  Board, keselamatan, dan kuis",
]
def replace_paragraph_text(shape, old, new):
    for paragraph in shape.text_frame.paragraphs:
        if old not in paragraph.text:
            continue
        replacement = paragraph.text.replace(old, new)
        if paragraph.runs:
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = replacement


for shape in agenda.shapes:
    if not shape.has_text_frame:
        continue
    for old, new in (
        ("00–10  Orientasi dan alur mesin", agenda_labels[0]),
        ("10–25  Jalankan simulasi dan jelajahi file", agenda_labels[1]),
        ("25–45  Praktik coding produk dan slot", agenda_labels[2]),
        ("45–65  Pahami protokol serial dan sketch", agenda_labels[3]),
        ("65–82  Praktik dan pengujian", agenda_labels[4]),
        ("82–90  Kuis, rangkuman, refleksi", agenda_labels[5]),
    ):
        replace_paragraph_text(shape, old, new)
    replace_paragraph_text(
        shape,
        "Jika waktu terbatas, ringkas penjelasan API pembayaran; "
        "tetap lakukan simulasi dan bahas keselamatan.",
        "Bagian board adalah demonstrasi instruktur; mahasiswa berlatih "
        "dengan simulasi dan perangkat tiruan.",
    )

for page_number, slide in enumerate(deck.slides, start=1):
    heading = next(
        (
            shape.text
            for shape in slide.shapes
            if shape.has_text_frame
            and ("PEMULA · LANGKAH" in shape.text
                 or "DEMONSTRASI INSTRUKTUR" in shape.text)
        ),
        "",
    )
    is_beginner = "PEMULA · LANGKAH" in heading
    badge_number = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        value = shape.text_frame.text.strip()
        if len(value) != 2 or not value.isdigit():
            continue
        left = shape.left / 914400
        top = shape.top / 914400
        if is_beginner and left < 2.0 and top < 6.8:
            badge_number += 1
            value = f"{badge_number:02d}"
        else:
            value = f"{page_number:02d}"
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = value
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(9 if left > 11 else 11)
        paragraph.font.color.rgb = (
            RGBColor(255, 255, 255) if left < 2.0
            else RGBColor(92, 103, 114)
        )
        if left > 11:
            paragraph.alignment = PP_ALIGN.RIGHT

deck.save(PATH)
print("Reordered guided slides and restored per-slide step numbers.")
