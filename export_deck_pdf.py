"""Export the workshop PowerPoint deck to a matching PDF."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.colors import Color, white
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


BASE = Path(r"D:\workshop-sukabumi\vending")
PPTX_PATH = BASE / "mini_vending_qris_workshop_2_hours.pptx"
PDF_PATH = BASE / "mini_vending_qris_workshop_2_hours.pdf"
FONT_PATH = Path(r"C:\Windows\Fonts")

pdfmetrics.registerFont(TTFont("Arial", str(FONT_PATH / "arial.ttf")))
pdfmetrics.registerFont(TTFont("Arial-Bold", str(FONT_PATH / "arialbd.ttf")))
pdfmetrics.registerFont(TTFont("Consolas", str(FONT_PATH / "consola.ttf")))

deck = Presentation(PPTX_PATH)
slide_width = deck.slide_width / 12700
slide_height = deck.slide_height / 12700
pdf = Canvas(str(PDF_PATH), pagesize=(slide_width, slide_height), pageCompression=1)
pdf.setTitle("Workshop on Building a Mini Vending Machine Integrated with QRIS Payments")
pdf.setAuthor("Ajang Rahmat, S.Kom")


def color_from_rgb(rgb):
    return Color(rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)


def get_run_color(run):
    if run is None:
        return Color(0.12, 0.15, 0.18)
    try:
        rgb = run.font.color.rgb
        if rgb is not None:
            return color_from_rgb(rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return Color(0.12, 0.15, 0.18)


def get_paragraph_color(paragraph, run):
    try:
        if run is not None:
            rgb = run.font.color.rgb
            if rgb is not None:
                return color_from_rgb(rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        rgb = paragraph.font.color.rgb
        if rgb is not None:
            return color_from_rgb(rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return Color(0.12, 0.15, 0.18)


def wrap_text(value, font_name, font_size, maximum_width):
    if not value:
        return [""]
    words = value.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if pdfmetrics.stringWidth(candidate, font_name, font_size) <= maximum_width:
            current = candidate
        elif current:
            lines.append(current)
            current = word
        else:
            lines.append(word)
    if current:
        lines.append(current)
    return lines


for slide in deck.slides:
    pdf.setFillColor(white)
    pdf.rect(0, 0, slide_width, slide_height, fill=1, stroke=0)

    for shape in slide.shapes:
        x = shape.left / 12700
        y_top = shape.top / 12700
        width = shape.width / 12700
        height = shape.height / 12700
        y_bottom = slide_height - y_top - height

        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                rgb = shape.fill.fore_color.rgb
                pdf.setFillColor(color_from_rgb(rgb))
                pdf.rect(x, y_bottom, width, height, fill=1, stroke=0)
            except (AttributeError, TypeError, ValueError):
                pass

        if not shape.has_text_frame:
            continue

        frame = shape.text_frame
        margin_left = frame.margin_left / 12700
        margin_right = frame.margin_right / 12700
        x_text = x + margin_left
        line_width = max(1, width - margin_left - margin_right)
        cursor_y = slide_height - y_top - frame.margin_top / 12700

        for paragraph in frame.paragraphs:
            run = paragraph.runs[0] if paragraph.runs else None
            paragraph_font = paragraph.font
            font_size = (
                run.font.size.pt if run and run.font.size
                else paragraph_font.size.pt if paragraph_font.size
                else 18
            )
            requested_font = (
                run.font.name if run and run.font.name
                else paragraph_font.name or "Arial"
            )
            bold = (
                run.font.bold if run and run.font.bold is not None
                else paragraph_font.bold
            )
            if requested_font.lower() == "consolas":
                font_name = "Consolas"
            else:
                font_name = "Arial-Bold" if bold else "Arial"
            pdf.setFont(font_name, font_size)
            pdf.setFillColor(get_paragraph_color(paragraph, run))

            lines = []
            for raw_line in paragraph.text.splitlines() or [""]:
                if font_name == "Consolas":
                    lines.append(raw_line)
                else:
                    lines.extend(
                        wrap_text(raw_line, font_name, font_size, line_width)
                    )

            leading = font_size * 1.22
            paragraph_bottom = slide_height - y_top - height - 2
            for line in lines:
                cursor_y -= leading
                if cursor_y < paragraph_bottom:
                    break
                string_width = pdfmetrics.stringWidth(line, font_name, font_size)
                if paragraph.alignment == 2:
                    draw_x = x_text + max(0, (line_width - string_width) / 2)
                elif paragraph.alignment == 3:
                    draw_x = x_text + max(0, line_width - string_width)
                else:
                    draw_x = x_text
                pdf.drawString(draw_x, cursor_y, line)

            if paragraph.space_after:
                cursor_y -= paragraph.space_after.pt

    pdf.showPage()

pdf.save()
print(f"Exported {len(deck.slides)} slides to {PDF_PATH}")
