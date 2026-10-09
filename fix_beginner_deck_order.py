"""Put the appended beginner slides in the correct teaching sequence."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PATH = r"D:\workshop-sukabumi\vending\materi_praktik_vending_90_menit.pptx"
deck = Presentation(PATH)
slide_ids = list(deck.slides._sldIdLst)
if len(slide_ids) != 27:
    raise ValueError(f"Expected 27 slides; found {len(slide_ids)}.")

ordered = (
    slide_ids[:22]
    + [slide_ids[26]]
    + [slide_ids[22]]
    + [slide_ids[25]]
    + slide_ids[23:25]
)
id_list = deck.slides._sldIdLst
for item in list(id_list):
    id_list.remove(item)
for item in ordered:
    id_list.append(item)

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 39, 49)
BLUE = RGBColor(39, 101, 161)
TEAL = RGBColor(35, 125, 126)
AMBER = RGBColor(167, 111, 43)
GRAY = RGBColor(92, 103, 114)
PALE_BLUE = RGBColor(238, 245, 251)
PALE_TEAL = RGBColor(237, 246, 244)
PALE_AMBER = RGBColor(249, 244, 235)
RULE = RGBColor(215, 221, 226)


def clear_slide(slide):
    tree = slide.shapes._spTree
    for child in list(tree)[2:]:
        tree.remove(child)


def add_text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
             align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    for index, row in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = row
        paragraph.alignment = align
        paragraph.space_after = Pt(2)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return shape


def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


# Rebuild the agenda as six brief, readable blocks.
agenda = deck.slides[2]
clear_slide(agenda)
add_text(agenda, "02 · TAHAPAN", 0.72, 0.33, 10.5, 0.23,
         size=10, color=BLUE, bold=True)
add_rect(agenda, 0.72, 0.78, 0.055, 0.47, BLUE)
add_text(agenda, "Peta praktik 90 menit", 0.92, 0.7, 11.6, 0.58,
         size=27, bold=True)
add_text(agenda, "Ikuti langkahnya berurutan; gunakan simulasi sebelum membahas board.",
         0.94, 1.37, 11.5, 0.38, size=13, color=GRAY)

agenda_items = [
    ("00–08", "Kenalan dengan proyek dan target hari ini", BLUE, PALE_BLUE),
    ("08–18", "Buka folder, siapkan terminal, jalankan aplikasi", TEAL, PALE_TEAL),
    ("18–28", "Coba tombol, pesanan, total, dan stok", AMBER, PALE_AMBER),
    ("28–48", "Ubah satu harga, jalankan ulang, cek hasil", BLUE, PALE_BLUE),
    ("48–73", "Telusuri kode dan uji serial tiruan", TEAL, PALE_TEAL),
    ("73–90", "Lihat sketch, keselamatan, tanya jawab", AMBER, PALE_AMBER),
]
for index, (minutes, task, accent, fill) in enumerate(agenda_items):
    column = index % 2
    row = index // 2
    x = 0.9 + column * 6.05
    y = 1.95 + row * 1.25
    add_rect(agenda, x, y, 5.68, 0.98, fill)
    add_rect(agenda, x, y, 0.06, 0.98, accent)
    add_text(agenda, minutes, x + 0.2, y + 0.17, 1.18, 0.31,
             size=14, color=accent, bold=True)
    add_text(agenda, task, x + 1.48, y + 0.15, 3.95, 0.62,
             size=14, bold=True)
add_text(agenda, "Catatan: bagian board adalah demonstrasi instruktur; tidak perlu transaksi live.",
         0.94, 6.02, 11.5, 0.36, size=13, color=GRAY)
add_rect(agenda, 0.72, 7.08, 11.9, 0.012, RULE)
add_text(agenda, "PRAKTIK VENDING · 90 MENIT", 0.74, 7.15, 7, 0.18,
         size=8, color=GRAY)
add_text(agenda, "03", 12.02, 7.13, 0.58, 0.2,
         size=9, color=GRAY, align=PP_ALIGN.RIGHT)

# Make all footer numbers match their actual order; restore instructional
# badges to 01, 02, ... within their own beginner slide.
for page_number, slide in enumerate(deck.slides, start=1):
    heading = next(
        (
            shape.text
            for shape in slide.shapes
            if shape.has_text_frame
            and ("PEMULA · LANGKAH" in shape.text
                 or "DEMONSTRASI INSTRUKTUR" in shape.text)
        ),
        "",
    )
    is_beginner = "PEMULA · LANGKAH" in heading
    badge_number = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        value = shape.text_frame.text.strip()
        if len(value) != 2 or not value.isdigit():
            continue
        left = shape.left / 914400
        top = shape.top / 914400
        if is_beginner and left < 2.0 and top < 6.8:
            badge_number += 1
            replacement = f"{badge_number:02d}"
            color = WHITE
            size = 11
        elif left > 11.0:
            replacement = f"{page_number:02d}"
            color = GRAY
            size = 9
        else:
            continue
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = replacement
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        if left > 11.0:
            paragraph.alignment = PP_ALIGN.RIGHT

# An empty PowerShell environment value is removed, so keep one space to
# override .env while the application treats the key as empty after strip().
for slide in deck.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text == '$env:SUMOPOD_API_KEY=""':
                    run.text = '$env:SUMOPOD_API_KEY=" "'
                elif run.text == (
                    "Key kosong hanya berlaku di terminal ini; "
                    "tidak menghapus atau mengubah .env."
                ):
                    run.text = (
                        "Satu spasi mencegah kunci dari .env terbaca; "
                        "aplikasi tetap menganggapnya kosong."
                    )

deck.save(PATH)
print("Corrected agenda, teaching order, and beginner step numbers.")
