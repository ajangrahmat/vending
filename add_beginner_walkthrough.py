"""Add guided, beginner-friendly setup and practice slides to the deck."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


DECK_PATH = r"D:\workshop-sukabumi\vending\materi_praktik_vending_90_menit.pptx"
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 39, 49)
MUTED = RGBColor(92, 103, 114)
BLUE = RGBColor(39, 101, 161)
TEAL = RGBColor(35, 125, 126)
ORANGE = RGBColor(167, 111, 43)
PALE_BLUE = RGBColor(238, 245, 251)
PALE_TEAL = RGBColor(237, 246, 244)
PALE_ORANGE = RGBColor(249, 244, 235)
PALE_GRAY = RGBColor(244, 246, 248)
RULE = RGBColor(215, 221, 226)

deck = Presentation(DECK_PATH)
width = deck.slide_width / 914400


def add_text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(4)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return shape


def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_card(slide, x, y, w, h, heading, body, accent=BLUE,
             fill=PALE_BLUE, body_size=14):
    add_rect(slide, x, y, w, h, fill)
    add_rect(slide, x, y, 0.06, h, accent)
    add_text(slide, heading, x + 0.2, y + 0.14, w - 0.4, 0.26,
             size=12, color=accent, bold=True)
    add_text(slide, body, x + 0.2, y + 0.48, w - 0.4, h - 0.55,
             size=body_size)


def add_step_slide(section, title, subtitle=None):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, section.upper(), 0.72, 0.33, 11.8, 0.23,
             size=10, color=BLUE, bold=True)
    add_rect(slide, 0.72, 0.78, 0.055, 0.45, BLUE)
    add_text(slide, title, 0.93, 0.71, 11.55, 0.58,
             size=26, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.95, 1.36, 11.35, 0.38,
                 size=13, color=MUTED)
    add_rect(slide, 0.72, 7.08, 11.9, 0.012, RULE)
    add_text(slide, "PRAKTIK VENDING · 90 MENIT", 0.74, 7.15, 7, 0.18,
             size=8, color=MUTED)
    add_text(slide, f"{len(deck.slides):02d}", 12.02, 7.13, 0.58, 0.2,
             size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_instructions(slide, steps, y=1.95, line=0.65):
    accents = [BLUE, TEAL, ORANGE]
    for index, step in enumerate(steps):
        top = y + index * line
        color = accents[index % len(accents)]
        add_rect(slide, 0.97, top + 0.04, 0.38, 0.34, color)
        add_text(slide, f"{index + 1:02d}", 1.01, top + 0.1, 0.3, 0.18,
                 size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, step, 1.56, top + 0.01, 10.65, line - 0.04,
                 size=16)


def add_code(slide, source, x=1.0, y=2.0, w=11.25, h=3.0, size=16):
    add_rect(slide, x, y, w, h, PALE_GRAY)
    add_rect(slide, x, y, 0.055, h, BLUE)
    box = slide.shapes.add_textbox(
        Inches(x + 0.2), Inches(y + 0.14), Inches(w - 0.4), Inches(h - 0.28)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    for index, line in enumerate(source.strip("\n").splitlines()):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.space_after = Pt(3)
        p.font.name = "Consolas"
        p.font.size = Pt(size)
        p.font.color.rgb = INK


def replace_text_on_slide(slide, old, new):
    for shape in slide.shapes:
        if shape.has_text_frame and old in shape.text:
            shape.text_frame.text = shape.text.replace(old, new)


# Retiming existing agenda for the guided path.
agenda = deck.slides[2]
for old, new in [
    ("00–10  Orientasi dan alur mesin", "00–08  Orientasi dan target belajar"),
    ("10–25  Jalankan simulasi dan jelajahi file", "08–18  Buka proyek, siapkan terminal, jalankan aplikasi"),
    ("25–45  Praktik coding produk dan slot", "18–28  Coba antarmuka dalam simulasi"),
    ("45–65  Pahami protokol serial dan sketch", "28–48  Ubah data produk, jalankan ulang, periksa"),
    ("65–82  Praktik dan pengujian", "48–73  Telusuri kode, protokol, dan tes serial tiruan"),
    ("82–90  Kuis, rangkuman, refleksi", "73–90  Bahas board, troubleshooting, exit ticket"),
]:
    replace_text_on_slide(agenda, old, new)
for shape in agenda.shapes:
    if shape.has_text_frame and "Jika waktu terbatas" in shape.text:
        shape.text_frame.text = (
            "Bagian board adalah demonstrasi instruktur; mahasiswa berlatih "
            "dengan simulasi dan perangkat tiruan."
        )

new_slides = []

slide = add_step_slide(
    "Pemula · langkah 1 dari 7",
    "Buka folder proyek di VS Code",
    "Tujuannya: terminal harus berada di folder yang memiliki main.py.",
)
add_instructions(slide, [
    "Buka VS Code dari menu Start.",
    "Pilih File → Open Folder…",
    "Pilih folder proyek vending, lalu pilih Select Folder.",
    "Jika ada pertanyaan Workspace Trust, percayai hanya folder proyek yang dikenal.",
    "Di Explorer sebelah kiri, pastikan main.py dan requirements.txt terlihat.",
], line=0.68)
add_card(slide, 1.0, 5.65, 11.2, 0.82, "JIKA FOLDER BELUM ADA",
         "Instruktur membagikan salinan proyek. Jangan mengedit folder .venv atau file .env.",
         ORANGE, PALE_ORANGE, 13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 2 dari 7",
    "Buka terminal dan siapkan mode simulasi",
    "Lakukan ini di terminal PowerShell baru di dalam VS Code.",
)
add_code(slide, """py --version
py -m pip install -r requirements.txt
$env:SUMOPOD_API_KEY=" "
py main.py""", y=2.0, h=2.8, size=18)
add_card(slide, 1.0, 5.08, 5.45, 1.12, "APA YANG TERJADI?",
         "Perintah pertama mengecek Python. Kedua memasang library proyek. Perintah ketiga memaksa simulasi.",
         BLUE, PALE_BLUE, 13)
add_card(slide, 6.85, 5.08, 5.45, 1.12, "AMAN UNTUK LATIHAN",
         "Satu spasi mencegah kunci dari .env terbaca; aplikasi tetap menganggapnya kosong.",
         TEAL, PALE_TEAL, 13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 3 dari 7",
    "Pastikan aplikasi benar-benar simulasi",
    "Jangan lanjutkan sebelum label pembayaran menunjukkan mode simulasi.",
)
add_instructions(slide, [
    "Tunggu jendela Mini Vending Machine terbuka.",
    "Lihat panel Status Sistem dan temukan tulisan Mode simulasi.",
    "Pilih satu produk; amati total dan stok yang berubah.",
    "Klik tombol minus untuk mengembalikan produk ke pesanan semula.",
    "Klik Kosongkan pesanan sebelum melanjutkan.",
], line=0.69)
add_card(slide, 1.0, 5.74, 11.2, 0.76, "JIKA TIDAK MELIHAT MODE SIMULASI",
         "Tutup aplikasi. Buka terminal PowerShell baru, ulangi perintah simulasi pada slide sebelumnya.",
         ORANGE, PALE_ORANGE, 13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 4 dari 7",
    "Edit satu harga dengan aman",
    "Gunakan VS Code. Jangan mengubah nama field atau nomor slot pada latihan ini.",
)
add_instructions(slide, [
    "Tekan Ctrl+P, ketik inventory.py, lalu tekan Enter.",
    "Tekan Ctrl+F, cari Air Mineral.",
    "Temukan baris 'price': 3000.",
    "Ganti 3000 menjadi 3500 saja.",
    "Tekan Ctrl+S. Titik putih pada tab file harus hilang.",
], y=1.92, line=0.57)
add_code(slide, """'name': "Air Mineral",
'price': 3500,  # hanya angka harga yang diubah
'stock': 15,
'image': "air_mineral.png",
'slot': 1,""", y=5.02, h=1.15, size=13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 5 dari 7",
    "Jalankan ulang dan buktikan perubahan",
    "Perubahan kode baru terbaca saat aplikasi dimulai ulang.",
)
add_instructions(slide, [
    "Kembali ke terminal tempat aplikasi sedang berjalan.",
    "Tekan Ctrl+C untuk menutup aplikasi.",
    "Jalankan py main.py lagi dalam terminal simulasi yang sama.",
    "Pastikan tulisan Mode simulasi masih terlihat.",
    "Pilih Air Mineral; cek harga kartu dan total.",
], y=1.94, line=0.63)
add_card(slide, 1.0, 5.58, 5.43, 0.9, "HASIL YANG DIHARAPKAN",
         "Harga berubah menjadi Rp 3.500 dan total mengikuti harga baru.",
         TEAL, PALE_TEAL, 13)
add_card(slide, 6.82, 5.58, 5.43, 0.9, "PULIHKAN DATA",
         "Ubah kembali ke 3000, simpan, tutup, lalu jalankan ulang.",
         BLUE, PALE_BLUE, 13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 6 dari 7",
    "Ikuti kode dengan Ctrl+Shift+F",
    "Cari nama fungsi, lalu telusuri satu pemanggilan pada satu waktu.",
)
add_instructions(slide, [
    "Tekan Ctrl+Shift+F untuk mencari di semua file proyek.",
    "Cari select_product; buka hasil dari main.py.",
    "Temukan add_to_cart — pesanan ditambahkan dan stok diperiksa di sana.",
    "Lihat update_products dan update_order — tampilan kemudian diperbarui.",
    "Cari dispense_product — alur sesudah pembayaran berhasil berlanjut ke dispenser.",
], y=1.92, line=0.68)
add_card(slide, 1.0, 5.68, 11.2, 0.8, "KATA KUNCI PENTING",
         "inventory.py = data produk    ·    cart.py = aturan keranjang    ·    serial_controller.py = pesan ke Arduino",
         BLUE, PALE_BLUE, 13)
new_slides.append(slide)

slide = add_step_slide(
    "Pemula · langkah 7 dari 7",
    "Tes dulu tanpa Arduino dan motor",
    "FakeSerial adalah perangkat tiruan untuk latihan protokol.",
)
add_code(slide, """py test_serial_controller.py
py test_main_gui.py""", y=2.02, h=1.55, size=18)
add_instructions(slide, [
    "Jalankan perintah pertama; tunggu pesan Semua uji serial lulus.",
    "Jalankan perintah kedua; pastikan tes GUI selesai dengan SEMUA LULUS.",
    "Baca test_serial_controller.py untuk melihat READY, PING/PONG, dan DONE.",
    "Ingat: tes tiruan tidak menguji kabel, relay, motor, atau produk nyata.",
], y=4.0, line=0.5)
new_slides.append(slide)

slide = add_step_slide(
    "Demonstrasi instruktur · board fisik",
    "Port COM dan hardware: hanya dengan pengawasan",
    "Mahasiswa tidak perlu melakukan transaksi live atau mengaktifkan motor sendiri.",
)
add_instructions(slide, [
    "Instruktur memastikan board dan pin sesuai dengan sketch.",
    "Hubungkan board; cek nomor COM di Device Manager.",
    "Tutup Serial Monitor sebelum aplikasi memakai port yang sama.",
    "Cocokkan relay aktif-LOW/HIGH, catu daya, dan durasi motor.",
    "Jangan jalankan pembayaran live untuk latihan kelas.",
], y=1.96, line=0.65)
add_card(slide, 1.0, 5.58, 11.2, 0.94, "BATASAN PENTING",
         "DONE berarti output motor selesai berjalan; belum membuktikan sensor mendeteksi produk jatuh.",
         ORANGE, PALE_ORANGE, 14)
new_slides.append(slide)

# Put guided setup after the project map, editing after the data example,
# and trace/testing immediately after their related material.
old_ids = list(deck.slides._sldIdLst)
new_ids = list(old_ids[-len(new_slides):])
kept_ids = old_ids[:-len(new_slides)]

layout = (
    kept_ids[:5]
    + new_ids[:3]
    + kept_ids[5:6]
    + new_ids[3:5]
    + kept_ids[6:7]
    + new_ids[5:6]
    + kept_ids[7:]
    + new_ids[6:]
)
slide_list = deck.slides._sldIdLst
for element in list(slide_list):
    slide_list.remove(element)
for element in layout:
    slide_list.append(element)

# Rewrite displayed slide numbers after inserting the beginner steps.
for index, slide in enumerate(deck.slides, start=1):
    for shape in slide.shapes:
        if (
            shape.has_text_frame
            and shape.text_frame.text.isdigit()
            and len(shape.text_frame.text) == 2
        ):
            p = shape.text_frame.paragraphs[0]
            p.text = f"{index:02d}"
            p.font.name = "Aptos"
            p.font.size = Pt(9)
            p.font.color.rgb = MUTED
            p.alignment = PP_ALIGN.RIGHT

deck.core_properties.title = "Praktik Vending · Panduan Langkah demi Langkah"
deck.core_properties.subject = "Workshop 90 menit untuk mahasiswa pemula"
deck.save(DECK_PATH)
print(f"Added 8 guided slides; deck now has {len(deck.slides)} slides.")
